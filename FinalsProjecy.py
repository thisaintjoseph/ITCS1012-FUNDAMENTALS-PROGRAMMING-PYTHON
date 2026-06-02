import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
import datetime
import openpyxl as op

#               ROOM RESERVATION SYSTEM

# ------------------ Room Rates ------------------
room_rates = {
    1: ("Single", 1000),
    2: ("Double", 1800),
    3: ("Suite", 3000),
    4: ("VIP", 5000)
}

# ------------------ Functions ------------------
def display():
    workbook = op.load_workbook("Afable_Database.xlsx")
    sheet = workbook.active
    table.delete(*table.get_children())
    for row in sheet.iter_rows(min_row=2, values_only=True):
        table.insert("", tk.END, values=row)

def generate_id():
    workbook = op.load_workbook("Afable_Database.xlsx")
    sheet = workbook.active
    return sheet.max_row

def compute_cost(room_number, nights):
    room_name, rate = room_rates.get(room_number, ("Unknown", 0))
    return room_name, rate * nights

def update_price_label(event=None):

    try:
        room_number = int(room_combo.get())
        room_name, rate = room_rates[room_number]
        price_label.config(text=f"Rate: ₱{rate} ({room_name})")
    except:
        price_label.config(text="Rate: -")

def add_record():
    name = name_entry.get().strip()
    room_number = room_combo.get().strip()
    date = date_entry.get_date().strftime("%Y-%m-%d")
    nights = nights_entry.get().strip()

    if not name or not room_number or not date or not nights:
        messagebox.showerror("Error", "All fields are required!")
        return

    try:
        room_number = int(room_number)
        nights = int(nights)
    except ValueError:
        messagebox.showerror("Error", "Room and Nights must be numbers!")
        return

    room_name, cost = compute_cost(room_number, nights)

    workbook = op.load_workbook("Afable_Database.xlsx")
    sheet = workbook.active
    new_id = generate_id()
    sheet.append([new_id, name, room_name, date, nights, cost])
    workbook.save("Afable_Database.xlsx")
    display()
    messagebox.showinfo("Success", f"Reservation added! Room: {room_name}, Cost: ₱{cost}")

def select_record(event):
    selected = table.focus()
    values = table.item(selected, "values")
    if values:
        
        
        name_entry.delete(0, tk.END)
        room_combo.set("")
        date_entry.delete(0, tk.END)
        nights_entry.delete(0, tk.END)
        
        name_entry.insert(0, values[1])

       
        for num, (room_name, _) in room_rates.items():
            if room_name == values[2]:
                room_combo.set(num)
                update_price_label()
                break

        try:
            date_obj = datetime.datetime.strptime(values[3], "%Y-%m-%d").date()
            date_entry.set_date(date_obj)
        except ValueError:
            date_entry.set_date(values[3])
        
        nights_entry.insert(0, values[4])

def update_record():
    selected = table.focus()
    if not selected:
        messagebox.showerror("Error", "No record selected!")
        return

    id_val = table.item(selected, "values")[0]
    name = name_entry.get().strip()
    room_number = room_combo.get().strip()
    date = date_entry.get_date().strftime("%Y-%m-%d")
    nights = nights_entry.get().strip()

    if not name or not room_number or not date or not nights:
        messagebox.showerror("Error", "All fields are required!")
        return

    try:
        room_number = int(room_number)
        nights = int(nights)
    except ValueError:
        messagebox.showerror("Error", "Room and Nights must be numbers!")
        return

    room_name, cost = compute_cost(room_number, nights)

    workbook = op.load_workbook("Afable_Database.xlsx")
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == id_val:
            row[1].value = name
            row[2].value = room_name
            row[3].value = date
            row[4].value = nights
            row[5].value = cost
            break
    workbook.save("Afable_Database.xlsx")
    display()
    messagebox.showinfo("Success", f"Reservation updated! Room: {room_name}, Cost: ₱{cost}")

def delete_record():
    selected = table.focus()
    if not selected:
        messagebox.showerror("Error", "No record selected!")
        return

    id_val = table.item(selected, "values")[0]
    workbook = op.load_workbook("Afable_Database.xlsx")
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=2):
        if str(row[0].value) == id_val:
            sheet.delete_rows(row[0].row)
            break
    workbook.save("Afable_Database.xlsx")
    display()
    messagebox.showinfo("Success", "Reservation deleted successfully!")

# ------------------ GUI ------------------
window = tk.Tk()
window.title("Room Reservation System")
window.geometry("800x500")
window.resizable(True, True)
window.configure(bg="#0b1d3a")

for i in range(7):  # rows used
    window.grid_rowconfigure(i, weight=1)
for j in range(3):  # columns used
    window.grid_columnconfigure(j, weight=1)

tk.Label(window, text="Room Reservation System",
         font=("Arial", 20, "bold"),
         bg="#0b1d3a", fg="#ffd700").grid(row=0, column=0, columnspan=3, pady=15)

# Labels & Entries

tk.Label(window, text="Name", bg="#0b1d3a", fg="#00ced1",).grid(row=1, column=0, padx=10, pady=5)
name_entry = tk.Entry(window)
name_entry.grid(row=1, column=1, padx=10, pady=5)

tk.Label(window, text="Room (1=Single, 2=Double, 3=Suite, 4 = VIP)", bg="#0b1d3a", fg="#adff2f",).grid(row=2, column=0, padx=10, pady=5)
room_combo = ttk.Combobox(window, values=[1, 2, 3, 4], state="readonly")
room_combo.grid(row=2, column=1, padx=10, pady=5)
room_combo.bind("<<ComboboxSelected>>", update_price_label) 

price_label = tk.Label(window, text="Rate: -", fg="#ffd700", bg="#0b1d3a")
price_label.grid(row=2, column=2, padx=10, pady=5)

tk.Label(window, text="Date", bg="#0b1d3a", fg="#ff69b4",).grid(row=3, column=0, padx=10, pady=5)
date_entry = DateEntry(window, width=18, background="darkblue",foreground="white", borderwidth=2, date_pattern="yyyy-mm-dd", mindate=datetime.date.today())
date_entry.grid(row=3, column=1, padx=10, pady=5)

tk.Label(window, text="Nights", bg="#0b1d3a", fg="#1e90ff",).grid(row=4, column=0, padx=10, pady=5)
nights_entry = tk.Entry(window)
nights_entry.grid(row=4, column=1, padx=10, pady=5)


# Buttons
tk.Button(window, text="Add", command=add_record, bg="#32cd32", fg="white", font=("Arial", 10, "bold")).grid(row=5, column=0, pady=10)
tk.Button(window, text="Update", command=update_record, bg="#1e90ff", fg="white", font=("Arial", 10, "bold")).grid(row=5, column=1, pady=10)
tk.Button(window, text="Delete", command=delete_record, bg="#dc143c", fg="white", font=("Arial", 10, "bold")).grid(row=5, column=2, pady=10)

style = ttk.Style()
style.configure("Treeview", background="#e6f7ff", foreground="black", rowheight=25, fieldbackground="#e6f7ff")
style.map("Treeview", background=[("selected", "#ffd700")])

frame = tk.Frame(window)
frame.grid(row=6, column=0, columnspan=3, padx=10, pady=10, sticky="nsew")

scrollbar = tk.Scrollbar(frame, orient="vertical")
scrollbar.pack(side="right", fill="y")

columns = ("ID", "Name", "Room", "Date", "Nights", "Cost")
table = ttk.Treeview(window, columns=columns, show="headings")
for col in columns:
    table.heading(col, text=col)
    table.column(col, width=120)
table.grid(row=6, column=0, columnspan=3, padx=10, pady=10)

table.bind("<<TreeviewSelect>>", select_record)

display()
window.mainloop()
